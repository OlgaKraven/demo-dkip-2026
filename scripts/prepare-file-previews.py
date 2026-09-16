"""Read source workbooks into an HTML preview; never modify the workbooks."""
from pathlib import Path
from html import escape
from datetime import date, datetime
import json
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parent.parent
previews = {}
for path in sorted((root / 'materials/practice/2026/assets').rglob('*.xlsx')):
    book = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    sheets = []
    for ws in book:
        used = [c for row in ws for c in row if c.value is not None]
        if not used:
            continue
        rows = max(c.row for c in used)
        cols = max(c.column for c in used)
        starts, covered = {}, set()
        for m in ws.merged_cells.ranges:
            if ws.cell(m.min_row, m.min_col).value is None:
                continue
            rows, cols = max(rows, m.max_row), max(cols, m.max_col)
            starts[(m.min_row, m.min_col)] = (m.max_row-m.min_row+1, m.max_col-m.min_col+1)
            covered.update((r,c) for r in range(m.min_row,m.max_row+1) for c in range(m.min_col,m.max_col+1) if (r,c)!=(m.min_row,m.min_col))
        html = ['<table><thead><tr><th></th>']
        html += ['<th>'+get_column_letter(c)+'</th>' for c in range(1,cols+1)]
        html.append('</tr></thead><tbody>')
        for r in range(1,rows+1):
            html.append('<tr><th>'+str(r)+'</th>')
            for c in range(1,cols+1):
                if (r,c) in covered:
                    continue
                cell=ws.cell(r,c)
                value=cached[ws.title].cell(r,c).value if cell.data_type=='f' else cell.value
                if value is None and cell.data_type=='f':
                    value=cell.value
                if isinstance(value,(datetime,date)):
                    value=value.strftime('%d.%m.%Y')
                if isinstance(value,float) and value.is_integer():
                    value=int(value)
                rowspan,colspan=starts.get((r,c),(1,1))
                css=' class="cell-bold"' if cell.font.bold else ''
                html.append(f'<td title="{cell.coordinate}" rowspan="{rowspan}" colspan="{colspan}"{css}>'+escape('' if value is None else str(value))+'</td>')
            html.append('</tr>')
        html.append('</tbody></table>')
        sheets.append({'name':ws.title,'html':''.join(html),'sourceCells':len(used)})
    key=path.relative_to(root/'materials').as_posix()
    previews[key]={'title':path.name,'sheets':sheets}
(root/'content/file-previews.json').write_text(json.dumps(previews,ensure_ascii=False),encoding='utf-8')
print(json.dumps({k:sum(s['sourceCells'] for s in v['sheets']) for k,v in previews.items()},ensure_ascii=False))
